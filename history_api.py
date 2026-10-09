"""Read-only AddLiveTag price/commission history integration; never stores API keys."""
import io
import math
import re
from datetime import date, datetime

import requests
from flask import request, jsonify, send_file
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

HISTORY_URL = 'https://data.addlivetag.com/price-tracking/history.php'


def num(v):
    if v is None or v == '' or isinstance(v, bool):
        return None
    try:
        value = float(v)
        return value if math.isfinite(value) else None
    except (ValueError, TypeError):
        return None


def clean_points(block, kind):
    rows = []
    for p in block.get('points', []):
        if not isinstance(p, dict):
            continue
        try:
            d = date.fromisoformat(str(p.get('date', ''))).isoformat()
        except ValueError:
            continue
        row = {'date': d, 'suspect': bool(p.get('suspect')), 'changed': bool(p.get('changed'))}
        fields = ['price', 'originalPrice', 'discountPercent'] if kind == 'price' else [
            'sellerRatePercent', 'shopeeRatePercent', 'totalRatePercent',
            'commission', 'sellerComFinal', 'shopeeComFinal', 'priceSnapshot']
        row.update({k: num(p.get(k)) for k in fields})
        if kind == 'commission':
            for field, decimal in [('sellerRatePercent', 'sellerRate'), ('shopeeRatePercent', 'shopeeRate')]:
                if row[field] is None and num(p.get(decimal)) is not None:
                    row[field] = num(p[decimal]) * 100
            if row['totalRatePercent'] is None and row['sellerRatePercent'] is not None and row['shopeeRatePercent'] is not None:
                row['totalRatePercent'] = row['sellerRatePercent'] + row['shopeeRatePercent']
            row['isCapped'] = bool(p.get('isCapped'))
            recorded = p.get('recorded') if isinstance(p.get('recorded'), dict) else {}
            row['recorded'] = {k: num(recorded.get(k)) for k in ['sellerRatePercent', 'shopeeRatePercent', 'totalRatePercent', 'commission']}
        else:
            row['currency'] = str(p.get('currency') or 'VND')
            row['recordedTime'] = str(p.get('recordedTime') or '')
        rows.append(row)
    return sorted(rows, key=lambda p: p['date'])


def normalize_history(item, item_id, data, settings):
    price = item.get('price') if isinstance(item.get('price'), dict) else {}
    com = item.get('commission') if isinstance(item.get('commission'), dict) else {}
    normalized = data.get('normalized') is True or item.get('normalized') is True or com.get('normalized') is True
    messages = [str(x) for x in [data.get('notice'), data.get('warning'), item.get('message'), item.get('warning'), com.get('warning')] if x]
    if not normalized:
        messages.append('Chưa xác nhận chuẩn hóa HH theo tài khoản. Tỷ lệ sàn lưu theo tài khoản thu thập từng ngày có thể tạo biến động giả; hãy khai HH sàn và trần của bạn.')
    if settings['base_rate'] is None or settings['cap'] is None:
        messages.append('Chưa khai đủ HH sàn và trần. Không coi tiền/% HH lịch sử là số liệu chính xác cho tài khoản của bạn.')
    limits = data.get('limits') if isinstance(data.get('limits'), dict) else {}
    truncated = bool(limits.get('truncated'))
    if truncated:
        messages.append('API đã cắt bớt điểm lịch sử. Hãy giảm số ngày hoặc chỉ tải một sản phẩm; biểu đồ và file này chưa đầy đủ.')
    prices = clean_points(price, 'price')
    commissions = clean_points(com, 'commission')
    if any(p['suspect'] for p in commissions):
        messages.append('Có điểm HH được nguồn đánh dấu nghi ngờ (suspect). Tool giữ nguyên số, không tự sửa hay loại bỏ.')
    status = item.get('status', 'error')
    if status == 'success' and not prices and not commissions:
        status = 'no_data'
    stats = {}
    for key, block in [('priceStats', price), ('commissionStats', com)]:
        raw = block.get('stats') if isinstance(block.get('stats'), dict) else {}
        allowed = ['min', 'max', 'avg', 'first', 'last', 'changeCount', 'dayCount', 'change', 'changePercent', 'minTotalRatePercent', 'maxTotalRatePercent', 'firstTotalRatePercent', 'lastTotalRatePercent', 'minCommission', 'maxCommission', 'sellerRateChangeCount', 'shopeeRateRecordedChangeCount']
        stats[key] = {k: num(raw.get(k)) for k in allowed}
    return {
        'itemId': str(item_id), 'status': status, 'prices': prices, 'commissions': commissions,
        **stats, 'normalized': normalized, 'truncated': truncated,
        'warnings': list(dict.fromkeys(messages)), 'range': data.get('range') or {},
        'changesOnly': data.get('changesOnly', True), 'settings': settings,
        'checkedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
    }


def register_history(app, safe_cell):
    @app.post('/api/history')
    def history():
        body = request.get_json(silent=True) or {}
        key = body.get('apiKey', '')
        if not isinstance(key, str) or not key.strip() or len(key) > 512 or not key.isascii() or any(ord(c) < 32 for c in key):
            return jsonify(error='Nhập API Key hợp lệ từ addlivetag.com/tool/api-key.php.'), 400
        key = key.strip()
        ids = body.get('itemIds')
        if not isinstance(ids, list) or not 1 <= len(ids) <= 50 or any(not re.fullmatch(r'[1-9]\d{0,19}', str(i)) for i in ids):
            return jsonify(error='Lịch sử nhận tối đa 50 Item ID hợp lệ mỗi lô.'), 400
        days = body.get('days', 90)
        if isinstance(days, bool) or not isinstance(days, int) or not 1 <= days <= 730:
            return jsonify(error='Khoảng lịch sử phải từ 1 đến 730 ngày.'), 400
        payload = {'item_ids': [str(i) for i in ids], 'type': 'both', 'days': days, 'changes_only': 1, 'no_product': 1}
        settings = {'days': days, 'base_rate': None, 'cap': None}
        for field, maximum in [('base_rate', 100), ('cap', 100000000)]:
            v = body.get(field)
            if v is not None and v != '':
                n = num(v)
                if n is None or not 0 <= n <= maximum:
                    return jsonify(error=f'{field} không hợp lệ.'), 400
                settings[field] = n
                payload[field] = n / 100 if field == 'base_rate' else n
        try:
            response = requests.post(HISTORY_URL, json=payload, headers={'X-API-Key': key, 'Accept': 'application/json'}, timeout=(12, 100), allow_redirects=False)
            if response.status_code in (401, 403):
                return jsonify(error='API lịch sử từ chối truy cập. Kiểm tra API Key và quyền tài khoản.'), response.status_code
            if response.status_code == 429:
                retry = response.headers.get('Retry-After', '60')
                return jsonify(error='API lịch sử đạt hạn mức. Đang chờ để thử lại.', retryAfter=max(60, int(retry) if retry.isdigit() else 60)), 429
            if response.status_code != 200:
                return jsonify(error=f'API lịch sử trả HTTP {response.status_code}.'), 502
            try:
                data = response.json()
            except ValueError:
                return jsonify(error='API lịch sử không trả JSON hợp lệ.'), 502
            if not isinstance(data, dict) or data.get('status') != 'success' or not isinstance(data.get('items'), list):
                return jsonify(error='Phản hồi lịch sử không đúng tài liệu.'), 502
            mapping = {str(x.get('itemId') or x.get('input')): x for x in data['items'] if isinstance(x, dict)}
            results = []
            for item_id in ids:
                item = mapping.get(str(item_id), {'status': 'error', 'message': 'API không trả lịch sử cho Item ID này.'})
                row = normalize_history(item, item_id, data, settings)
                row['warnings'] = [w.replace(key, '[đã ẩn]') for w in row['warnings']]
                results.append(row)
            return jsonify(results=results, limits=data.get('limits') or {})
        except requests.Timeout:
            return jsonify(error='API lịch sử phản hồi quá lâu. Dữ liệu hiện tại của sản phẩm vẫn được giữ.'), 504
        except requests.RequestException:
            return jsonify(error='Không kết nối được API lịch sử. Dữ liệu hiện tại của sản phẩm vẫn được giữ.'), 502

    @app.post('/api/history/export')
    def history_export():
        body = request.get_json(silent=True) or {}
        histories = body.get('histories', [])
        if not isinstance(histories, list) or not 1 <= len(histories) <= 500 or any(not isinstance(h, dict) for h in histories):
            return jsonify(error='Xuất từ 1 đến 500 sản phẩm lịch sử mỗi lần.'), 400
        total = 0
        for h in histories:
            for field in ('prices', 'commissions'):
                points = h.get(field, [])
                if not isinstance(points, list) or any(not isinstance(p, dict) for p in points):
                    return jsonify(error='Điểm lịch sử không hợp lệ.'), 400
                total += len(points)
        if total > 100000:
            return jsonify(error='Tối đa 100.000 điểm mỗi file. Hãy lọc ít sản phẩm hoặc giảm số ngày.'), 400
        wb = Workbook()
        summary = wb.active
        summary.title = 'Tổng quan'
        summary.append(['Item ID', 'Sản phẩm', 'Trạng thái', 'Từ', 'Đến', 'HH chuẩn hóa', 'HH sàn đã khai (%)', 'Trần đã khai (VND)', 'Bị cắt dữ liệu', 'Cảnh báo', 'Thời gian tải'])
        price_sheet = wb.create_sheet('Lịch sử giá')
        price_sheet.append(['Item ID', 'Ngày', 'Giá (VND)', 'Giá gốc (VND)', 'Giảm giá (%)', 'Điểm thay đổi', 'Thời gian ghi'])
        com_sheet = wb.create_sheet('Lịch sử hoa hồng')
        com_sheet.append(['Item ID', 'Ngày', 'HH sàn (%)', 'HH Xtra (%)', 'Tổng HH (%)', 'Tiền HH (VND)', 'Giá snapshot (VND)', 'Nghi ngờ', 'Áp trần', 'Tỷ lệ sàn gốc (%)', 'Tổng tỷ lệ gốc (%)', 'Tiền HH gốc (VND)'])
        for h in histories:
            settings = h.get('settings') or {}
            rg = h.get('range') or {}
            warnings = h.get('warnings') or []
            summary.append([safe_cell(v) for v in [h.get('itemId'), h.get('name'), h.get('status'), rg.get('from'), rg.get('to'), h.get('normalized'), settings.get('base_rate'), settings.get('cap'), h.get('truncated'), ' | '.join(map(str, warnings)), h.get('checkedAt')]])
            for p in h.get('prices', []):
                price_sheet.append([safe_cell(v) for v in [h.get('itemId'), p.get('date'), p.get('price'), p.get('originalPrice'), p.get('discountPercent'), p.get('changed'), p.get('recordedTime')]])
            for p in h.get('commissions', []):
                rec = p.get('recorded') or {}
                com_sheet.append([safe_cell(v) for v in [h.get('itemId'), p.get('date'), p.get('shopeeRatePercent'), p.get('sellerRatePercent'), p.get('totalRatePercent'), p.get('commission'), p.get('priceSnapshot'), p.get('suspect'), p.get('isCapped'), rec.get('shopeeRatePercent'), rec.get('totalRatePercent'), rec.get('commission')]])
        for ws in wb:
            ws.freeze_panes = 'C2'
            ws.auto_filter.ref = ws.dimensions
            for c in ws[1]:
                c.fill = PatternFill('solid', fgColor='124B44')
                c.font = Font(bold=True, color='FFFFFF')
                ws.column_dimensions[c.column_letter].width = 26
        notes = wb.create_sheet('Lưu ý')
        for text in [
            'DỮ LIỆU GIẢ LẬP, KHÔNG PHẢI SỐ LIỆU THỰC.' if body.get('demo') else 'Nguồn: AddLiveTag Price & Commission History API (bên thứ ba).',
            'Lịch sử theo NGÀY, giữ lần ghi cuối trong ngày; không phản ánh mọi thay đổi trong ngày.',
            'Chỉ lấy mốc thay đổi. Ô thiếu dữ liệu là chưa biết, không phải 0.',
            'HH sàn có thể thay đổi giả do tài khoản thu thập khác tier. Kiểm tra cột chuẩn hóa, cấu hình và cảnh báo.',
            'Điểm nghi ngờ được giữ nguyên; không tự sửa. Bị cắt dữ liệu nghĩa là file chưa đầy đủ.',
            'Không tính lượt bán theo ngày từ API này. Lượt bán hiện tại là historical sales từ Product Data Batch.',
            'Chỉ dùng học tập/nghiên cứu/nội bộ phi thương mại. Không bảo đảm hoa hồng thanh toán thực tế.',
        ]:
            notes.append([text])
        notes.column_dimensions['A'].width = 140
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', download_name='LinkScope-history' + ('-DEMO' if body.get('demo') else '') + '.xlsx')
