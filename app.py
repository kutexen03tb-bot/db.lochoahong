import csv
import io
import math
import os
import re
from datetime import datetime
from urllib.parse import urlparse, parse_qs

import requests
from flask import Flask, render_template, request, jsonify, send_file
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 12 * 1024 * 1024
API_URL = 'https://data.addlivetag.com/product-data/product-data-batch.php'
MAX_ROWS = 10000


def number(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    try:
        n = float(value)
        return n if math.isfinite(n) and n >= 0 else None
    except (TypeError, ValueError):
        return None


def extract_id(raw):
    raw = str(raw).strip().strip('"').strip("'")
    if re.fullmatch(r'[1-9]\d{0,19}', raw):
        return raw, None
    try:
        u = urlparse(raw if '://' in raw else 'https://' + raw)
        host = (u.hostname or '').lower()
    except ValueError:
        return None, 'URL không hợp lệ'
    if u.scheme not in ('http', 'https') or u.username or u.password:
        return None, 'URL không hợp lệ'
    if host in ('s.shopee.vn', 'shp.ee', 'vn.shp.ee') or host.endswith('.shp.ee'):
        return None, 'Link rút gọn: hãy mở trên trình duyệt và lấy link sản phẩm đầy đủ'
    if host not in ('shopee.vn', 'www.shopee.vn', 'm.shopee.vn'):
        return None, 'Cần link sản phẩm shopee.vn hoặc Item ID'
    match = re.search(r'-i\.(\d+)\.([1-9]\d*)(?:/|$)', u.path)
    if not match:
        match = re.search(r'/(?:product|opaanlp)/(\d+)/([1-9]\d*)(?:/|$)', u.path)
    if match:
        return match[2], None
    q = parse_qs(u.query)
    for key in ('item_id', 'itemId'):
        value = q.get(key, [''])[0]
        if re.fullmatch(r'[1-9]\d{0,19}', value):
            return value, None
    return None, 'Không tìm thấy Item ID trong link'


def normalize(product, item_id, original):
    p = product.get('productInfo') or {}
    status = product.get('status', 'error')
    source = product.get('dataSource') or p.get('dataSource') or ''
    reason = product.get('reason') or ''
    usable = status in ('success', 'stale') and source != 'fallback' and bool(p)
    verified = usable and reason != 'commission_unverified'
    warning = product.get('warning') or product.get('message') or reason
    def n(field):
        return number(p.get(field)) if usable else None
    def rate(percent, decimal):
        if not verified:
            return None
        v = number(p.get(percent))
        if v is not None:
            return v
        v = number(p.get(decimal))
        return v * 100 if v is not None else None
    sr = rate('sellerRatePercent', 'sellerRate')
    br = rate('shopeeRatePercent', 'shopeeRate')
    total = number(p.get('totalRatePercent')) if verified else None
    if total is None and sr is not None and br is not None:
        total = sr + br
    if status == 'success' and not usable:
        status = 'incomplete'
        warning = warning or 'API không cung cấp dữ liệu đã xác minh'
    if usable and (n('price') is None or total is None):
        warning = warning or 'Thiếu giá hoặc tỷ lệ hoa hồng; không quy đổi tiền HH thành tỷ lệ công bố'
        if status == 'success':
            status = 'incomplete'
    price = n('price')
    commission = n('commission') if verified else None
    url = p.get('productLink') or p.get('originLink') or original
    if not isinstance(url, str) or extract_id(url)[1] or not url.startswith(('https://', 'http://')):
        url = ''
    return {
        'itemId': str(item_id), 'input': original, 'name': p.get('productName') or '',
        'shop': p.get('shopName') or '', 'url': url, 'price': price,
        'sellerRate': sr, 'shopeeRate': br, 'totalRate': total,
        'commission': commission, 'sellerCommission': n('sellerComFinal') if verified else None,
        'shopeeCommission': n('shopeeComFinal') if verified else None,
        'effectiveRate': commission / price * 100 if commission is not None and price and price > 0 else None,
        'capped': bool(p.get('isCapped') or p.get('isLimitCap')), 'cap': n('capRaw'),
        'source': source, 'status': status, 'warning': str(warning or ''),
        'updated': p.get('lastUpdate') or '', 'rateSource': p.get('shopeeRateSource') or '',
        'sales': n('sales'), 'rating': n('rating'),
        'checkedAt': datetime.now().astimezone().isoformat(timespec='seconds'),
    }


@app.after_request
def headers(response):
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; object-src 'none'; base-uri 'self'; form-action 'self'"
    return response


@app.errorhandler(413)
def large_file(e):
    return jsonify(error='File quá lớn. Tối đa 12 MB.'), 413


@app.get('/')
def index():
    return render_template('index.html')


@app.post('/api/analyze')
def analyze():
    body = request.get_json(silent=True) or {}
    values = body.get('values', [])
    if not isinstance(values, list) or len(values) > MAX_ROWS:
        return jsonify(error='Tối đa 10.000 dòng mỗi lần nhập.'), 400
    seen, items, errors, duplicates = set(), [], [], 0
    for idx, value in enumerate(values, 1):
        if not isinstance(value, (str, int, float)) or not str(value).strip():
            continue
        original = str(value).strip()
        if len(original) > 4096:
            errors.append({'line': idx, 'input': original[:100], 'error': 'Link quá dài'})
            continue
        item_id, error = extract_id(original)
        if error:
            errors.append({'line': idx, 'input': original, 'error': error})
        elif item_id in seen:
            duplicates += 1
        else:
            seen.add(item_id)
            items.append({'itemId': item_id, 'input': original})
    return jsonify(items=items, errors=errors, duplicates=duplicates)


@app.post('/api/import')
def import_file():
    f = request.files.get('file')
    if not f:
        return jsonify(error='Chưa chọn file.'), 400
    name = (f.filename or '').lower()
    try:
        if name.endswith('.xlsx'):
            # Reject unusually expanded archives before parsing user-supplied workbooks.
            import zipfile
            buffer = io.BytesIO(f.read())
            with zipfile.ZipFile(buffer) as archive:
                if sum(z.file_size for z in archive.infolist()) > 60 * 1024 * 1024:
                    raise ValueError('File Excel giải nén quá lớn.')
            buffer.seek(0)
            wb = load_workbook(buffer, data_only=True, read_only=False)
            ws = wb.worksheets[0]
            if ws.max_row > MAX_ROWS + 1 or ws.max_column > 100:
                wb.close()
                raise ValueError('Tối đa 10.000 dòng dữ liệu và 100 cột; chỉ đọc sheet đầu tiên.')
            rows = [[str(c.hyperlink.target if c.hyperlink and c.hyperlink.target else (c.value if c.value is not None else '')) for c in row] for row in ws.iter_rows()]
            sheet = ws.title
            wb.close()
        elif name.endswith('.csv'):
            text = f.read().decode('utf-8-sig')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text), dialect)
            rows = []
            for row in reader:
                if len(row) > 100 or len(rows) > MAX_ROWS:
                    raise ValueError('Tối đa 10.000 dòng dữ liệu và 100 cột.')
                rows.append(row)
            sheet = 'CSV'
        else:
            return jsonify(error='Chỉ hỗ trợ .xlsx hoặc .csv (UTF-8).'), 400
        rows = [r for r in rows if any(str(c).strip() for c in r)]
        if not rows:
            raise ValueError('File không có dữ liệu.')
        return jsonify(rows=rows, sheet=sheet, filename=f.filename)
    except UnicodeDecodeError:
        return jsonify(error='Hãy lưu CSV với mã hóa UTF-8 hoặc dùng file .xlsx.'), 400
    except ValueError as e:
        return jsonify(error=str(e)), 400
    except Exception:
        return jsonify(error='Không đọc được file. Hãy kiểm tra định dạng và bỏ mật khẩu bảo vệ nếu có.'), 400


@app.post('/api/check')
def check():
    body = request.get_json(silent=True) or {}
    key = body.get('apiKey', '')
    if not isinstance(key, str) or not key.strip() or len(key) > 512 or re.search(r'[\r\n]', key):
        return jsonify(error='Nhập API Key hợp lệ từ addlivetag.com.'), 400
    key = key.strip()
    if not key.isascii() or any(ord(c) < 32 for c in key):
        return jsonify(error='API Key chứa ký tự không hợp lệ.'), 400
    items = body.get('items', [])
    if not isinstance(items, list) or not 1 <= len(items) <= 100:
        return jsonify(error='Mỗi lô phải có từ 1 đến 100 sản phẩm.'), 400
    if any(not isinstance(p, dict) or not re.fullmatch(r'[1-9]\d{0,19}', str(p.get('itemId', ''))) for p in items):
        return jsonify(error='Item ID không hợp lệ.'), 400
    payload = {'item_ids': [str(p['itemId']) for p in items], 'max_api': 20}
    if body.get('cacheOnly'):
        payload.update(cache_only=1, max_api=0)
    for field, maximum in [('base_rate', 100), ('cap', 100000000)]:
        value = body.get(field)
        if value is not None and value != '':
            v = number(value)
            if v is None or v > maximum:
                return jsonify(error=f'{field} không hợp lệ.'), 400
            # UI uses percent units; convert to fraction to avoid the provider's ambiguous 0..1 input.
            payload[field] = v / 100 if field == 'base_rate' else v
    try:
        upstream = requests.post(API_URL, json=payload, headers={'X-API-Key': key, 'Accept': 'application/json'}, timeout=(12, 100), allow_redirects=False)
        if upstream.status_code in (401, 403):
            return jsonify(error='API từ chối truy cập. Kiểm tra key, quyền và trạng thái tài khoản.'), upstream.status_code
        if upstream.status_code == 429:
            retry = upstream.headers.get('Retry-After', '60')
            retry = int(retry) if retry.isdigit() else 60
            return jsonify(error='Đạt hạn mức API. Tool sẽ chờ trước khi thử lại.', retryAfter=max(60, retry)), 429
        if not upstream.ok or 300 <= upstream.status_code < 400:
            return jsonify(error=f'Nhà cung cấp trả HTTP {upstream.status_code}. Vui lòng thử lại sau.'), 502
        try:
            data = upstream.json()
        except ValueError:
            return jsonify(error='API không trả JSON hợp lệ.'), 502
        if not isinstance(data, dict) or data.get('status') != 'success' or not isinstance(data.get('products'), list):
            return jsonify(error='API trả lỗi hoặc cấu trúc phản hồi không đúng tài liệu batch.'), 502
        mapped = {}
        for p in data['products']:
            if isinstance(p, dict):
                pid = p.get('itemId') or (p.get('productInfo') or {}).get('itemId') or p.get('input')
                mapped[str(pid)] = p
        results = []
        for item in items:
            item_id = str(item['itemId'])
            p = mapped.get(item_id, {'status': 'error', 'message': 'API không trả dòng tương ứng với Item ID này'})
            row = normalize(p, item_id, str(item.get('input', '')))
            row['warning'] = row['warning'].replace(key, '[đã ẩn]')
            results.append(row)
        return jsonify(results=results, limits=data.get('limits') or {}, warning=str(data.get('warning') or '').replace(key, '[đã ẩn]'))
    except requests.Timeout:
        return jsonify(error='API phản hồi quá lâu. Lô này chưa được xác nhận; có thể thử lại sau.'), 504
    except requests.RequestException:
        return jsonify(error='Không kết nối được API nhà cung cấp.'), 502


EXPORT_FIELDS = [
    ('itemId', 'Item ID'), ('input', 'Link đầu vào'), ('name', 'Sản phẩm'), ('shop', 'Shop'),
    ('price', 'Giá bán (VND)'), ('shopeeRate', 'HH sàn (%)'), ('sellerRate', 'HH Xtra (%)'),
    ('totalRate', 'Tổng HH công bố (%)'), ('commission', 'Tiền HH theo API (VND)'),
    ('effectiveRate', 'Tỷ lệ tiền HH / giá (%)'), ('shopeeCommission', 'Tiền HH sàn (VND)'),
    ('sellerCommission', 'Tiền HH Xtra (VND)'), ('capped', 'Bị áp trần'), ('source', 'Nguồn'),
    ('status', 'Trạng thái'), ('updated', 'Cập nhật tại nguồn'), ('checkedAt', 'Thời gian kiểm tra'),
    ('rateSource', 'Nguồn tỷ lệ'), ('warning', 'Lưu ý'), ('url', 'Link sản phẩm'),
    ('sales', 'Tổng đã bán'), ('rating', 'Đánh giá')]


def safe_cell(value):
    if isinstance(value, (list, dict)):
        value = str(value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


@app.post('/api/export')
def export():
    body = request.get_json(silent=True) or {}
    rows = body.get('rows', [])
    if not isinstance(rows, list) or len(rows) > MAX_ROWS or any(not isinstance(x, dict) for x in rows):
        return jsonify(error='Danh sách xuất không hợp lệ.'), 400
    demo = bool(body.get('demo'))
    filename = 'shopee-' + ('DEMO-' if demo else '') + datetime.now().strftime('%Y%m%d-%H%M')
    if body.get('format') == 'csv':
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow([v for k, v in EXPORT_FIELDS])
        w.writerows([[safe_cell(row.get(k)) for k, v in EXPORT_FIELDS] for row in rows])
        return send_file(io.BytesIO(out.getvalue().encode('utf-8-sig')), mimetype='text/csv', as_attachment=True, download_name=filename + '.csv')
    wb = Workbook()
    ws = wb.active
    ws.title = 'DEMO - Giả lập' if demo else 'Kết quả đã lọc'
    ws.append([v for k, v in EXPORT_FIELDS])
    for row in rows:
        ws.append([safe_cell(row.get(k)) for k, v in EXPORT_FIELDS])
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='124B44')
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 24 if letter not in ('B', 'C', 'S', 'T') else 55
    for row in ws.iter_rows(min_row=2):
        for idx in (4, 5, 6, 7, 8, 9, 10, 11):
            row[idx].number_format = '#,##0.00'
    note = wb.create_sheet('Lưu ý')
    for text in [
        'DỮ LIỆU GIẢ LẬP, KHÔNG PHẢI GIÁ/HOA HỒNG THỰC.' if demo else 'Nguồn: addlivetag.com, API bên thứ ba; không phải cam kết hoa hồng được thanh toán.',
        'Tài liệu giới hạn mục đích học tập/nghiên cứu/vận hành nội bộ phi thương mại.',
        'Giá là trường price của API; không bảo đảm là giá thanh toán sau voucher hoặc đúng mọi biến thể.',
        'Tổng % HH dùng các trường tỷ lệ của API; ô trống là chưa biết, không phải 0%.',
        'Tiền HH và tỷ lệ tiền HH/giá có thể khác tỷ lệ công bố vì trần, user rate và thuế.',
        'Dữ liệu có thể cache 3 giờ hoặc cũ hơn khi có lỗi. Kiểm tra Nguồn, Trạng thái, Lưu ý.',
        'Tất cả trường % trong bảng là đơn vị phần trăm: giá trị 5.5 nghĩa là 5,5%.',
        'Thời gian cập nhật tại nguồn theo Asia/Ho_Chi_Minh, theo tài liệu nhà cung cấp.',
    ]:
        note.append([text])
    note.column_dimensions['A'].width = 135
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(stream, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=filename + '.xlsx')


from history_api import register_history
register_history(app, safe_cell)


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', '8000')), debug=False)
